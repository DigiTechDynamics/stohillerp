"""
Tabular exports in the format the user asked for (UAT #12: "Excel" and "PDF"
used to both download a CSV).

A report is given as rows: a list of lists of cells. The first row is the
title row; a row that follows an empty row (or the first data row) and has no
numbers is treated as a header. The same rows render as:

    csv   text/csv
    xlsx  a real Excel workbook (numbers stored as numbers)
    pdf   an A4 report with the company header (landscape when wide)

    export_response(rows, 'xlsx', 'trial-balance')
"""

import csv
import io
from decimal import Decimal, InvalidOperation

from django.http import HttpResponse

CONTENT_TYPES = {
    'csv': 'text/csv',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'pdf': 'application/pdf',
}
FORMAT_ALIASES = {'excel': 'xlsx', 'xls': 'xlsx'}


def normalise_format(value) -> str:
    value = (value or 'csv').lower()
    value = FORMAT_ALIASES.get(value, value)
    return value if value in CONTENT_TYPES else 'csv'


def _number(cell):
    """Decimal for numeric-looking cells, else None."""
    if isinstance(cell, bool) or cell is None:
        return None
    if isinstance(cell, (int, float, Decimal)):
        return Decimal(str(cell))
    text = str(cell).strip().replace(',', '')
    if not text or not any(ch.isdigit() for ch in text):
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _header_rows(rows):
    """Indexes of rows that look like column headers."""
    headers = set()
    for i, row in enumerate(rows):
        if i == 0 or not row or not any(str(c).strip() for c in row):
            continue
        previous_blank = i > 0 and not any(str(c).strip() for c in rows[i - 1])
        if (previous_blank or i == 1) and len(row) > 2 and all(_number(c) is None for c in row):
            headers.add(i)
    return headers


def to_csv(rows) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    # BOM so Excel opens UTF-8 CSVs (names with accents) correctly.
    return ('﻿' + buf.getvalue()).encode('utf-8')


def to_xlsx(rows, title='Report') -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (title or 'Report')[:31]
    headers = _header_rows(rows)
    header_fill = PatternFill('solid', fgColor='1E293B')
    widths = {}
    for r, row in enumerate(rows, start=1):
        for c, cell in enumerate(row, start=1):
            number = _number(cell)
            value = number if number is not None else ('' if cell is None else str(cell))
            target = ws.cell(row=r, column=c, value=float(number) if number is not None else value)
            if number is not None:
                target.number_format = '#,##0.00' if number != number.to_integral_value() or '.' in str(cell) else '0'
            if r == 1:
                target.font = Font(bold=True, size=13)
            elif (r - 1) in headers:
                target.font = Font(bold=True, color='FFFFFF')
                target.fill = header_fill
            widths[c] = max(widths.get(c, 8), min(len(str(cell or '')) + 2, 60))
    for c, width in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rows, title='Report') -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from apps.finance.services import pdf_service as p

    width = max((len(r) for r in rows), default=1)
    pagesize = landscape(A4) if width > 5 else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, topMargin=p.MARGIN, bottomMargin=p.MARGIN,
                            leftMargin=p.MARGIN, rightMargin=p.MARGIN)
    s = p._base_styles()
    story = []
    heading = rows[0] if rows else [title]
    p._header_block(story, s, escape(str(heading[0] if heading else title)).upper(),
                    ref_label='', ref_value='', date_label='Generated')
    if len(heading) > 1:
        story.append(Paragraph(escape(' '.join(str(c) for c in heading[1:] if c)), s['body']))
        story.append(Spacer(1, 4 * mm))

    body = rows[1:]
    headers = {i - 1 for i in _header_rows(rows) if i >= 1}
    table_rows, styles = [], []
    for i, row in enumerate(body):
        cells = list(row) + [''] * (width - len(row))
        if not any(str(c).strip() for c in cells):
            table_rows.append([''] * width)
            continue
        is_header = i in headers
        is_section = not is_header and sum(1 for c in cells if str(c).strip()) == 1 and _number(cells[0]) is None
        rendered = []
        for c, cell in enumerate(cells):
            # The first column holds labels and account codes, never amounts.
            number = _number(cell) if c > 0 else None
            text = f'{number:,.2f}' if number is not None else escape('' if cell is None else str(cell))
            style = s['right_bold'] if (number is not None and is_header) else \
                s['right'] if number is not None else s['body_bold'] if (is_header or is_section) else s['body']
            rendered.append(Paragraph(text, style))
        table_rows.append(rendered)
        row_index = len(table_rows) - 1
        if is_header:
            styles += [('BACKGROUND', (0, row_index), (-1, row_index), p.TABLE_STRIPE),
                       ('LINEBELOW', (0, row_index), (-1, row_index), 0.75, p.BRAND_GOLD)]
        elif is_section:
            styles.append(('LINEBELOW', (0, row_index), (-1, row_index), 0.5, p.TABLE_STRIPE))
        else:
            label = ' '.join(str(c) for c in cells).lower()
            if label.startswith(('total', 'net ', 'closing')) or ' total' in label[:20]:
                styles.append(('LINEABOVE', (0, row_index), (-1, row_index), 0.75, p.BRAND_DARK))
    if table_rows:
        usable = pagesize[0] - 2 * p.MARGIN
        first = usable * (0.34 if width > 2 else 0.6)
        col_widths = [first] + [(usable - first) / max(width - 1, 1)] * (width - 1)
        table = Table(table_rows, colWidths=col_widths, repeatRows=0)
        table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ] + styles))
        story.append(table)
    story.extend(p._footer_paragraph(s, 'Report'))
    doc.build(story)
    return buf.getvalue()


def export_response(rows, fmt, filename, title=None) -> HttpResponse:
    fmt = normalise_format(fmt)
    title = title or (str(rows[0][0]) if rows and rows[0] else filename)
    content = {'csv': lambda: to_csv(rows), 'xlsx': lambda: to_xlsx(rows, title), 'pdf': lambda: to_pdf(rows, title)}[fmt]()
    response = HttpResponse(content, content_type=CONTENT_TYPES[fmt])
    response['Content-Disposition'] = f'attachment; filename="{filename}.{fmt}"'
    return response
